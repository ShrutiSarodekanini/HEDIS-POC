"""Render the synthetic clinical evidence documents for the HEDIS POC.

Reads scripts/seed/scenario.py (the single source of truth) and writes:
  data/seed/documents/<doc_id>.<ext>   one file per scenario document, in its own format
  data/seed/documents/manifest.json    metadata and ground truth text for every document
  data/seed/documents/samples/         extra intake test files plus manifest_samples.json
  data/seed/structured/                CSV and XLSX structured files plus structured_expected.json

Run from the repo root:  python scripts/seed/gen_documents.py [--verify]

Output is deterministic: every random draw comes from a numpy generator seeded per document,
and every embedded timestamp is fixed, so reruns produce byte-identical files.
All people, places and values are fictional.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import shutil
import sys
import zipfile
import zlib
from datetime import date, datetime, timezone
from email.message import EmailMessage
from email.utils import format_datetime
from functools import lru_cache
from pathlib import Path

import numpy as np
from docx import Document
from fpdf import FPDF
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pypdf import PdfReader, PdfWriter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import scenario as sc  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
DOCS_DIR = ROOT / "data" / "seed" / "documents"
SAMPLES_DIR = DOCS_DIR / "samples"
STRUCT_DIR = ROOT / "data" / "seed" / "structured"
FONTS = Path("C:/Windows/Fonts")

SEED = 20261012
FIXED_DT = datetime(2026, 10, 12, 8, 0, 0)  # embedded in PDF, DOCX and XLSX metadata
ZIP_DT = (2026, 10, 12, 8, 0, 0)

EXT = {"txt": "txt", "pdf": "pdf", "scan_pdf": "pdf", "photo_jpg": "jpg", "handwritten_jpg": "jpg",
       "docx": "docx", "tiff_fax": "tif", "png": "png"}
MIME = {"txt": "text/plain", "pdf": "application/pdf", "scan_pdf": "application/pdf",
        "photo_jpg": "image/jpeg", "handwritten_jpg": "image/jpeg",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "tiff_fax": "image/tiff", "png": "image/png"}

ORGS = {o["id"]: o for o in sc.ORGANIZATIONS}
PROVS = {p["id"]: p for p in sc.PROVIDERS}
MEMBERS = {m["member_id"]: m for m in sc.MEMBERS}
TIFF_COMPRESSION = "group4"  # Pillow with libtiff supports Group 4; no fallback needed here


# ---------------------------------------------------------------- small helpers

def rng_for(doc_id: str) -> np.random.Generator:
    """Per-document generator so each file is independent of generation order."""
    return np.random.default_rng(zlib.crc32(f"{SEED}:{doc_id}".encode()))


@lru_cache(maxsize=None)
def font(name: str, px: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(FONTS / name), px)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_bytes(path: Path, data: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)


def normalize_zip(data: bytes) -> bytes:
    """Rewrite a zip (DOCX/XLSX) with fixed member timestamps so the bytes are reproducible."""
    out = io.BytesIO()
    stamp = FIXED_DT.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            zi = zipfile.ZipInfo(info.filename, ZIP_DT)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            body = src.read(info.filename)
            if info.filename == "docProps/core.xml":  # openpyxl stamps "now" as modified on save
                body = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>" + stamp, body)
            dst.writestr(zi, body)
    return out.getvalue()


def wrap_lines(text: str, fnt: ImageFont.FreeTypeFont, max_w: float) -> list[str]:
    """Word-wrap text, keeping blank lines and explicit line breaks."""
    lines: list[str] = []
    for para in text.split("\n"):
        if not para.strip():
            lines.append("")
            continue
        cur = ""
        for word in para.split(" "):
            cand = word if not cur else f"{cur} {word}"
            if cur and fnt.getlength(cand) > max_w:
                lines.append(cur)
                cur = word
            else:
                cur = cand
        lines.append(cur)
    return lines


def org_lines(org_id: str) -> tuple[str, str]:
    org = ORGS[org_id]
    place = ", ".join(x for x in (org["city"], org["state"]) if x)
    return org["name"], place


# ---------------------------------------------------------------- page rendering (images)

def render_text_page(text: str, font_file: str, pt: float, dpi: int, size: tuple[int, int],
                     paper: tuple[int, int, int], ink: tuple[int, int, int] = (22, 22, 22),
                     table: dict | None = None, margin_in: float = 1.0, fit_min: int | None = None) -> Image.Image:
    """Draw plain text (or a gridded table) onto a blank page image.

    With fit_min set, the page is cropped to its content (but not below fit_min pixels tall), which
    gives the half-sheet crops people usually photograph.
    """
    img = Image.new("RGB", size, paper)
    d = ImageDraw.Draw(img)
    px = round(pt * dpi / 72)
    margin = round(margin_in * dpi)
    if table:
        return _fit(_render_table_page(img, d, table, dpi, margin, ink), fit_min, margin)
    f = font(font_file, px)
    y = margin
    for line in wrap_lines(text, f, size[0] - 2 * margin):
        if line:
            d.text((margin, y), line, font=f, fill=ink)
        y += round(px * 1.35)
    return _fit(img, fit_min, margin, y)


def _fit(img: Image.Image, fit_min: int | None, margin: int, bottom: int | None = None) -> Image.Image:
    if fit_min is None:
        return img
    bottom = img.height if bottom is None else bottom
    return img.crop((0, 0, img.width, min(img.height, max(fit_min, bottom + margin))))


def _render_table_page(img: Image.Image, d: ImageDraw.ImageDraw, table: dict, dpi: int, margin: int,
                       ink: tuple[int, int, int]) -> Image.Image:
    """Vitals flowsheet: bold title above a ruled grid with a shaded header row."""
    px = round(16 * dpi / 72)
    title_f, head_f, cell_f = font("arialbd.ttf", round(17 * dpi / 72)), font("arialbd.ttf", px), font("arial.ttf", px)
    y = margin
    for line in wrap_lines(table["title"], title_f, img.width - 2 * margin):
        d.text((margin, y), line, font=title_f, fill=ink)
        y += round(px * 1.5)
    y += round(px * 1.2)
    cols, rows = table["columns"], table["rows"]
    total_w = img.width - 2 * margin
    weights = [1.3, 1.0, 0.8, 1.2][: len(cols)] if len(cols) == 4 else [1.0] * len(cols)
    widths = [round(total_w * w / sum(weights)) for w in weights]
    row_h, pad = round(px * 2.3), round(px * 0.6)
    for r_idx, row in enumerate([cols] + rows):
        x = margin
        for c_idx, cell in enumerate(row):
            box = [x, y, x + widths[c_idx], y + row_h]
            d.rectangle(box, fill=(226, 230, 236) if r_idx == 0 else None, outline=ink, width=3)
            d.text((x + pad, y + (row_h - px) // 2 - 4), cell, font=head_f if r_idx == 0 else cell_f, fill=ink)
            x += widths[c_idx]
        y += row_h
    return img.crop((0, 0, img.width, y + margin))


def perspective_coeffs(dst: list[tuple[float, float]], src: list[tuple[float, float]]) -> list[float]:
    """Coefficients for Image.transform(PERSPECTIVE): output points dst map to input points src."""
    a, b = [], []
    for (x, y), (u, v) in zip(dst, src):
        a.append([x, y, 1, 0, 0, 0, -u * x, -u * y])
        b.append(u)
        a.append([0, 0, 0, x, y, 1, -v * x, -v * y])
        b.append(v)
    return np.linalg.solve(np.array(a, float), np.array(b, float)).tolist()


def lighting(h: int, w: int, rng: np.random.Generator, strength: float, blob: float) -> np.ndarray:
    """Uneven illumination: a linear gradient plus one soft shadow blob, as a (h, w) multiplier."""
    yy, xx = np.mgrid[0:h, 0:w].astype(np.float32)
    ang = rng.uniform(0, 2 * np.pi)
    proj = (xx / w - 0.5) * np.cos(ang) + (yy / h - 0.5) * np.sin(ang)
    light = 1 - strength * (proj + 0.7) / 1.4
    cx, cy, r = rng.uniform(0, w), rng.uniform(0, h), rng.uniform(0.3, 0.6) * max(w, h)
    light *= 1 - blob * np.exp(-(((xx - cx) ** 2 + (yy - cy) ** 2) / (2 * (r / 2) ** 2)))
    return light.astype(np.float32)


def desk_background(size: tuple[int, int], rng: np.random.Generator) -> Image.Image:
    """Dark wood-ish desk with soft low-frequency texture."""
    w, h = size
    low = rng.normal(0, 1, (max(h // 40, 2), max(w // 40, 2)))
    tex = np.asarray(Image.fromarray(np.clip(low * 40 + 128, 0, 255).astype(np.uint8))
                     .resize(size, Image.Resampling.BICUBIC), dtype=np.float32) / 255
    base = np.array([96, 76, 60], np.float32)
    arr = base[None, None, :] * (0.7 + 0.6 * tex)[:, :, None]
    return Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8))


def phone_photo(page: Image.Image, rng: np.random.Generator, quality: str = "normal", mild: bool = False) -> Image.Image:
    """Make a page image look like a phone photo: warp, uneven light, shadow, blur, noise."""
    pw, ph = page.size
    strength, blob, blur, noise, bright = (0.14, 0.12, 0.8, 3.0, 1.0) if mild else (0.30, 0.22, 1.0, 4.0, 1.0)
    if quality == "poor_light":
        strength, blob, blur, noise, bright = 0.45, 0.28, 1.2, 11.0, 0.64
    elif quality == "blurry":
        blur, noise = 3.0, 5.0
    lit = np.asarray(page, np.float32) * lighting(ph, pw, rng, strength, blob)[:, :, None]
    page = Image.fromarray(np.clip(lit, 0, 255).astype(np.uint8))
    mx, my = int(pw * 0.09), int(ph * 0.07)
    cw, ch = pw + 2 * mx, ph + 2 * my
    j = (0.012 if mild else 0.022) * max(pw, ph)
    corners = [(mx, my), (mx + pw, my), (mx + pw, my + ph), (mx, my + ph)]
    quad = [(x + rng.uniform(-j, j), y + rng.uniform(-j, j)) for x, y in corners]
    coeffs = perspective_coeffs(quad, [(0, 0), (pw, 0), (pw, ph), (0, ph)])
    warped = page.transform((cw, ch), Image.Transform.PERSPECTIVE, coeffs, Image.Resampling.BICUBIC)
    mask = Image.new("L", page.size, 255).transform((cw, ch), Image.Transform.PERSPECTIVE, coeffs,
                                                    Image.Resampling.BICUBIC)
    bg = desk_background((cw, ch), rng)
    shadow = ImageFilter.GaussianBlur(cw * 0.012)
    shade = mask.filter(shadow).point(lambda v: int(v * 0.55))
    bg.paste((20, 16, 12), (int(cw * 0.012), int(cw * 0.016)), shade.crop((0, 0, cw - int(cw * 0.012),
                                                                           ch - int(cw * 0.016))))
    bg.paste(warped, mask=mask.filter(ImageFilter.GaussianBlur(1.0)))
    bg = bg.filter(ImageFilter.GaussianBlur(blur))
    bg = bg.resize((round(cw * 0.8), round(ch * 0.8)), Image.Resampling.LANCZOS)
    arr = np.asarray(bg, np.float32) * bright
    arr += rng.normal(0, noise, arr.shape).astype(np.float32)
    if quality == "poor_light":
        arr *= np.array([1.0, 0.95, 0.85], np.float32)  # warm, dim indoor cast
    return Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8))


def save_jpeg(img: Image.Image, path: Path) -> None:
    img.save(path, "JPEG", quality=80)


def scan_image(text: str, rng: np.random.Generator, font_file: str = "times.ttf") -> Image.Image:
    """Letter page at 300 dpi, rotated 0.5 to 1.5 degrees on a light gray bed, with speckle noise."""
    page = render_text_page(text, font_file, 12, 300, (2550, 3300), (250, 248, 242))
    angle = float(rng.uniform(0.5, 1.5) * rng.choice([-1, 1]))
    page = page.rotate(angle, resample=Image.Resampling.BICUBIC, fillcolor=(226, 226, 226))
    gray = np.asarray(page.convert("L"), np.float32)
    gray += rng.normal(0, 3.0, gray.shape).astype(np.float32)
    img = Image.fromarray(np.clip(gray, 0, 255).astype(np.uint8))
    d = ImageDraw.Draw(img)
    for _ in range(1400):
        x, y, r = int(rng.integers(0, img.width)), int(rng.integers(0, img.height)), int(rng.integers(1, 4))
        shade = int(rng.integers(40, 150))
        d.ellipse([x - r, y - r, x + r, y + r], fill=shade)
    return img.filter(ImageFilter.GaussianBlur(0.6))


def handwritten_page(text: str, rng: np.random.Generator) -> Image.Image:
    """Lined note page with per-character jitter in size, rotation and baseline, in dark blue ink."""
    w, h, line_h, top, left = 1700, 1500, 80, 230, 230
    img = Image.new("RGB", (w, h), (251, 247, 233))
    d = ImageDraw.Draw(img)
    for y in range(top, h - 60, line_h):
        d.line([(0, y), (w, y)], fill=(176, 196, 226), width=2)
    d.line([(left - 40, 0), (left - 40, h)], fill=(214, 120, 120), width=3)
    base = 58
    lines = wrap_lines(text, font("Inkfree.ttf", base), w - left - 120)
    for i, line in enumerate(lines):
        baseline = top + i * line_h - 8 + float(rng.normal(0, 2))
        drift = float(rng.uniform(-3, 3))
        x = left + float(rng.uniform(0, 8))
        for k, ch in enumerate(line):
            size = max(int(base * rng.uniform(0.9, 1.1)), 30)
            fnt = font("Inkfree.ttf", size)
            adv = fnt.getlength(ch)
            if ch != " ":
                tile_w = int(size * 2.4)
                tile = Image.new("L", (tile_w, tile_w), 0)
                ImageDraw.Draw(tile).text((tile_w / 2 - adv / 2, tile_w / 2 + size * 0.3), ch, font=fnt,
                                          fill=255, anchor="ls")
                tile = tile.rotate(float(rng.normal(0, 3.0)), resample=Image.Resampling.BICUBIC)
                ink = tuple(int(c + rng.integers(-8, 9)) for c in (22, 34, 112))
                cy = baseline + drift * k / max(len(line), 1) + float(rng.normal(0, 1.8)) - size * 0.3
                img.paste(ink, (int(x + adv / 2 - tile_w / 2), int(cy - tile_w / 2)),
                          tile.point(lambda v: int(v * 0.93)))
            x += adv * float(rng.uniform(0.97, 1.05)) + 1.5
    return img


# ---------------------------------------------------------------- per-format renderers

def make_txt(doc: dict, path: Path) -> None:
    write_bytes(path, doc["text"].encode("utf-8"))


def make_pdf(doc: dict, path: Path) -> None:
    """Digital PDF: org letterhead then the text in Arial, with a real text layer."""
    name, place = org_lines(doc["org"])
    pdf = FPDF(format="Letter")
    pdf.set_creation_date(FIXED_DT)
    pdf.set_title(f"{doc['type']} {doc['doc_id']}")
    pdf.set_author("Synthetic HEDIS POC data")
    pdf.set_margins(20, 18, 20)
    pdf.add_font("Arial", "", str(FONTS / "arial.ttf"))
    pdf.add_font("Arial", "B", str(FONTS / "arialbd.ttf"))
    pdf.add_page()
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 8, name, new_x="LMARGIN", new_y="NEXT")
    if place:
        pdf.set_font("Arial", "", 9)
        pdf.set_text_color(90, 90, 90)
        pdf.cell(0, 5, place, new_x="LMARGIN", new_y="NEXT")
    pdf.set_draw_color(60, 60, 60)
    pdf.set_line_width(0.4)
    pdf.line(20, pdf.get_y() + 2, pdf.w - 20, pdf.get_y() + 2)
    pdf.ln(8)
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Arial", "", 11)
    for line in doc["text"].split("\n"):
        if line.strip():
            pdf.multi_cell(0, 5.6, line, new_x="LMARGIN", new_y="NEXT")
        else:
            pdf.ln(5.6)
    pdf.output(str(path))


def make_scan_pdf(doc: dict, path: Path) -> None:
    img = scan_image(doc["text"], rng_for(doc["doc_id"]), "times.ttf" if len(doc["doc_id"]) % 2 else "arial.ttf")
    img.save(path, "PDF", resolution=300.0, creationDate="D:20261012080000Z", modDate="D:20261012080000Z")


def make_photo_jpg(doc: dict, path: Path) -> None:
    rng = rng_for(doc["doc_id"])
    fname = "arial.ttf" if zlib.crc32(doc["doc_id"].encode()) % 2 else "times.ttf"
    page = render_text_page(doc["text"], fname, 15, 200, (1700, 2200), (248, 246, 240),
                            table=doc.get("table"), fit_min=900)
    save_jpeg(phone_photo(page, rng, doc.get("photo_quality") or "normal"), path)


def make_handwritten_jpg(doc: dict, path: Path) -> None:
    rng = rng_for(doc["doc_id"])
    save_jpeg(phone_photo(handwritten_page(doc["text"], rng), rng, "normal", mild=True), path)


def make_docx(doc: dict, path: Path) -> None:
    name, place = org_lines(doc["org"])
    d = Document()
    d.add_heading(name, level=1)
    if place:
        d.add_paragraph(place)
    for line in doc["text"].split("\n"):
        d.add_paragraph(line)
    props = d.core_properties
    props.author = props.last_modified_by = "Synthetic HEDIS POC data"
    props.created = props.modified = FIXED_DT
    props.title = f"{doc['type']} {doc['doc_id']}"
    props.revision = 1
    buf = io.BytesIO()
    d.save(buf)
    write_bytes(path, normalize_zip(buf.getvalue()))


def to_bitonal(gray: np.ndarray, rng: np.random.Generator) -> Image.Image:
    """Fax look: soft blur, noise, hard threshold, a few dropout streaks and specks."""
    img = Image.fromarray(np.clip(gray, 0, 255).astype(np.uint8)).filter(ImageFilter.GaussianBlur(0.7))
    arr = np.asarray(img, np.float32) + rng.normal(0, 9, gray.shape).astype(np.float32)
    bw = Image.fromarray(np.where(arr > 150, 255, 0).astype(np.uint8), "L")
    d = ImageDraw.Draw(bw)
    for _ in range(4):  # dropout streaks, kept below the text area so they do not strike words
        y, x0 = int(rng.integers(1750, bw.height)), int(rng.integers(0, bw.width // 2))
        d.line([(x0, y), (x0 + int(rng.integers(200, 700)), y)], fill=0, width=1)
    for _ in range(160):
        x, y = int(rng.integers(0, bw.width)), int(rng.integers(0, bw.height))
        d.point((x, y), fill=0)
    return bw.convert("1", dither=Image.Dither.NONE)


def make_tiff_fax(doc: dict, path: Path) -> None:
    """Two-page fax: cover sheet, then the letter. 1-bit, 204x196 dpi, Group 4."""
    rng = rng_for(doc["doc_id"])
    size = (1734, 2156)
    org = ORGS[doc["org"]]
    sender = org["name"].upper()
    m = re.search(r"To:\s*(.+?)(?:\s{2,}|$)", doc["text"].split("\n")[1] if "\n" in doc["text"] else "")
    pcp = MEMBERS.get(doc["member_id"] or "", {}).get("pcp_org")
    to_name = m.group(1) if m else (ORGS[pcp]["name"] if pcp else "Health Plan Medical Records")
    stamp = f"{doc['date']} 09:14   FROM: {sender}   P."

    def strip(img: Image.Image, page_no: int) -> None:
        ImageDraw.Draw(img).text((120, 60), f"{stamp}{page_no}/2", font=font("cour.ttf", 26), fill=(0, 0, 0))

    cover = Image.new("RGB", size, (255, 255, 255))
    d = ImageDraw.Draw(cover)
    d.text((150, 250), "FAX", font=font("arialbd.ttf", 150), fill=(0, 0, 0))
    d.line([(150, 440), (size[0] - 150, 440)], fill=(0, 0, 0), width=6)
    rows = [("From:", f"{org['name']}, {org['city']}, {org['state']}"), ("To:", to_name),
            ("Date:", doc["date"]), ("Pages:", "2 (including cover sheet)"), ("Re:", "Medical records request")]
    y = 520
    for label, value in rows:
        d.text((150, y), label, font=font("arialbd.ttf", 48), fill=(0, 0, 0))
        d.text((520, y), value, font=font("arial.ttf", 48), fill=(0, 0, 0))
        y += 110
    note = ("CONFIDENTIAL: This fax contains protected health information. If you received it in error, "
            "notify the sender and destroy all copies.")
    for i, line in enumerate(wrap_lines(note, font("arial.ttf", 34), size[0] - 300)):
        d.text((150, 1500 + i * 48), line, font=font("arial.ttf", 34), fill=(0, 0, 0))
    strip(cover, 1)
    letter = render_text_page(doc["text"], "arial.ttf", 13, 200, size, (255, 255, 255), margin_in=0.9)
    strip(letter, 2)
    pages = [to_bitonal(np.asarray(p.convert("L"), np.float32), rng) for p in (cover, letter)]
    pages[0].save(path, save_all=True, append_images=pages[1:], compression=TIFF_COMPRESSION, dpi=(204, 196))


def make_png(doc: dict, path: Path) -> None:
    """Member portal message rendered like a screenshot: top bar, sidebar, message card."""
    w = 1280
    f_title, f_bold, f_body, f_small = (font("segoeuib.ttf", 30), font("segoeuib.ttf", 20),
                                        font("segoeui.ttf", 22), font("segoeui.ttf", 18))
    subject, *body = doc["text"].split("\n")
    body_lines = wrap_lines("\n".join(body), f_body, w - 220 - 140 - 80)
    h = max(760, 64 + 300 + 36 * len(body_lines) + 190)
    img = Image.new("RGB", (w, h), (244, 246, 249))
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, w, 64], fill=(31, 78, 121))
    d.text((28, 16), "Member Portal", font=font("segoeuib.ttf", 28), fill=(255, 255, 255))
    d.text((w - 330, 22), "Secure Messages  |  Sign out", font=f_small, fill=(214, 226, 240))
    d.rectangle([0, 64, 220, h], fill=(233, 238, 245))
    for i, item in enumerate(["Dashboard", "Messages", "Uploads", "Claims", "Care gaps", "Settings"]):
        y = 100 + i * 52
        if item == "Messages":
            d.rectangle([0, y - 10, 220, y + 36], fill=(207, 221, 238))
            d.rectangle([0, y - 10, 5, y + 36], fill=(31, 78, 121))
        d.text((28, y), item, font=f_bold if item == "Messages" else f_body, fill=(40, 50, 65))
    card = [250, 96, w - 40, h - 40]
    d.rounded_rectangle(card, radius=10, fill=(255, 255, 255), outline=(208, 214, 224), width=2)
    d.text((290, 124), subject, font=f_title, fill=(25, 32, 45))
    d.text((290, 176), f"Received: {doc['date']}   |   Source: Member upload   |   Document {doc['doc_id']}",
           font=f_small, fill=(110, 118, 132))
    d.line([(290, 220), (w - 80, 220)], fill=(222, 226, 234), width=2)
    y = 248
    for line in body_lines:
        d.text((290, y), line, font=f_body, fill=(35, 42, 55))
        y += 36
    d.rounded_rectangle([290, h - 120, 410, h - 76], radius=8, fill=(31, 78, 121))
    d.text((322, h - 111), "Reply", font=f_bold, fill=(255, 255, 255))
    d.text((440, h - 109), "Attachments (0)", font=f_small, fill=(110, 118, 132))
    img.save(path, "PNG")


MAKERS = {"txt": make_txt, "pdf": make_pdf, "scan_pdf": make_scan_pdf, "photo_jpg": make_photo_jpg,
          "handwritten_jpg": make_handwritten_jpg, "docx": make_docx, "tiff_fax": make_tiff_fax, "png": make_png}


def page_count(fmt: str, path: Path) -> int:
    if EXT[fmt] == "pdf":
        return len(PdfReader(str(path)).pages)
    if fmt == "tiff_fax":
        with Image.open(path) as im:
            return getattr(im, "n_frames", 1)
    return 1


# ---------------------------------------------------------------- manifest, samples, structured

def build_documents() -> list[dict]:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    manifest = []
    for doc in sc.DOCUMENTS:
        fmt = doc["format"]
        path = DOCS_DIR / f"{doc['doc_id']}.{EXT[fmt]}"
        MAKERS[fmt](doc, path)
        quality = doc.get("photo_quality") or ("normal" if fmt in ("photo_jpg", "handwritten_jpg") else None)
        manifest.append({
            "doc_id": doc["doc_id"], "member_id": doc["member_id"], "file": path.name, "format": fmt,
            "mime_type": MIME[fmt], "type": doc["type"], "date": doc["date"],
            "source_org": ORGS[doc["org"]]["name"], "author": PROVS[doc["author"]]["name"] if doc["author"] else None,
            "handwritten": fmt == "handwritten_jpg", "photo_quality": quality,
            "duplicate_of": doc.get("duplicate_of"), "member_hint": doc.get("member_hint"),
            "page_count": page_count(fmt, path), "sha256": sha256(path), "ground_truth_text": doc["text"],
        })
    (DOCS_DIR / "manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return manifest


def build_samples() -> None:
    """Intake edge cases: mixed PDF, zip bundle, email, spoofed extension, byte-identical copy."""
    SAMPLES_DIR.mkdir(parents=True, exist_ok=True)
    entries = []

    def add(name: str, purpose: str, mime: str, sources: list[str], expected: str) -> None:
        p = SAMPLES_DIR / name
        entries.append({"file": name, "purpose": purpose, "mime_type": mime, "source_docs": sources,
                        "expected_behavior": expected, "sha256": sha256(p)})

    writer = PdfWriter()
    writer.add_page(PdfReader(str(DOCS_DIR / "D-1003.pdf")).pages[0])
    writer.add_page(PdfReader(str(DOCS_DIR / "D-202.pdf")).pages[0])
    writer.add_metadata({"/Title": "mixed_digital_and_scanned", "/Producer": "pypdf"})
    buf = io.BytesIO()
    writer.write(buf)
    write_bytes(SAMPLES_DIR / "mixed_digital_and_scanned.pdf", buf.getvalue())
    add("mixed_digital_and_scanned.pdf", "PDF whose page 1 has a text layer and page 2 is image-only; intake must "
        "route per page (text extraction for page 1, OCR for page 2).", "application/pdf", ["D-1003", "D-202"],
        "2 pages; page 1 text layer, page 2 needs OCR")

    zbuf = io.BytesIO()
    with zipfile.ZipFile(zbuf, "w", zipfile.ZIP_DEFLATED) as z:
        for name in ("D-1102.png", "D-502.jpg"):
            zi = zipfile.ZipInfo(name, ZIP_DT)
            zi.compress_type, zi.external_attr = zipfile.ZIP_DEFLATED, 0o644 << 16
            z.writestr(zi, (DOCS_DIR / name).read_bytes())
    write_bytes(SAMPLES_DIR / "portal_upload_bundle.zip", zbuf.getvalue())
    add("portal_upload_bundle.zip", "Member portal upload as a zip archive; intake must unpack and process each "
        "member file.", "application/zip", ["D-1102", "D-502"], "2 entries: D-1102.png and D-502.jpg")

    msg = EmailMessage()
    msg["From"], msg["To"] = "records@lab.example.org", "intake@plan.example.org"
    msg["Subject"] = "Lab result for member M-009"
    msg["Date"] = format_datetime(datetime(2026, 10, 12, 9, 30, tzinfo=timezone.utc))
    msg["Message-ID"] = "<d-901-sample@lab.example.org>"
    msg.set_content("Hello,\n\nPlease find the lab result for member M-009 attached.\n\nRecords Desk")
    msg.add_attachment((DOCS_DIR / "D-901.pdf").read_bytes(), maintype="application", subtype="pdf",
                       filename="D-901.pdf")
    msg.set_boundary("==HEDIS_POC_SAMPLE_BOUNDARY==")
    write_bytes(SAMPLES_DIR / "email_with_attachment.eml", msg.as_bytes())
    add("email_with_attachment.eml", "Plain MIME email with a PDF attachment; intake must extract the attachment.",
        "message/rfc822", ["D-901"], "1 attachment: D-901.pdf")

    write_bytes(SAMPLES_DIR / "spoofed_extension.pdf", (DOCS_DIR / "D-301.jpg").read_bytes())
    add("spoofed_extension.pdf", "JPEG bytes saved with a .pdf name; intake must detect the real type from content, "
        "not the extension.", "image/jpeg", ["D-301"], "Detected as image/jpeg despite .pdf extension")

    shutil.copyfile(DOCS_DIR / "D-1003.pdf", SAMPLES_DIR / "duplicate_copy_of_D-1003.pdf")
    add("duplicate_copy_of_D-1003.pdf", "Byte-identical copy of D-1003.pdf; intake must dedupe by content hash.",
        "application/pdf", ["D-1003"], "Same sha256 as D-1003.pdf; flagged as exact duplicate")

    entries.append({"file": None, "purpose": "Antivirus test: EICAR test file.", "mime_type": None, "source_docs": [],
                    "expected_behavior": "Quarantined or rejected by the scanner", "sha256": None,
                    "note": "create at test time"})
    (SAMPLES_DIR / "manifest_samples.json").write_text(json.dumps(entries, indent=2) + "\n", encoding="utf-8")


def build_structured() -> None:
    STRUCT_DIR.mkdir(parents=True, exist_ok=True)
    expected = []
    for sf in sc.STRUCTURED_FILES:
        path = STRUCT_DIR / sf["file"]
        if sf["format"] == "csv":
            with path.open("w", newline="", encoding="utf-8") as fh:
                w = csv.writer(fh, lineterminator="\n")
                w.writerow(sf["columns"])
                w.writerows(r["values"] for r in sf["rows"])
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "Vitals"
            ws.append(sf["columns"])
            for r in sf["rows"]:
                ws.append([_cell(c, v) for c, v in zip(sf["columns"], r["values"])])
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    if isinstance(c.value, date):
                        c.number_format = "yyyy-mm-dd"
            wb.properties.creator = wb.properties.lastModifiedBy = "Synthetic HEDIS POC data"
            wb.properties.created = wb.properties.modified = FIXED_DT
            buf = io.BytesIO()
            wb.save(buf)
            write_bytes(path, normalize_zip(buf.getvalue()))
        for i, r in enumerate(sf["rows"], start=1):
            expected.append({"file": sf["file"], "row": i, "expect": r["expect"], "reason": r.get("reason"),
                             "closes": r.get("closes"), "note": r.get("note")})
    (STRUCT_DIR / "structured_expected.json").write_text(json.dumps(expected, indent=2) + "\n", encoding="utf-8")


def _cell(column: str, value):
    """XLSX cell value: real Excel dates for date columns, empty cells for None and empty strings."""
    if value is None or value == "":
        return None
    if "Date" in column:
        return date.fromisoformat(value)
    return value


# ---------------------------------------------------------------- verification

def norm(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip().lower()


def verify() -> list[str]:
    """Check outputs against the scenario. Returns a list of failure messages."""
    fails: list[str] = []
    manifest = {m["doc_id"]: m for m in json.loads((DOCS_DIR / "manifest.json").read_text(encoding="utf-8"))}
    quotes: dict[str, list[str]] = {}
    for g in sc.GAPS:
        for item in g["expected"].get("find", []) + g["expected"].get("optional", []):
            quotes.setdefault(item["doc"], []).append(item["quote"])
    for doc in sc.DOCUMENTS:
        did, fmt = doc["doc_id"], doc["format"]
        path = DOCS_DIR / manifest[did]["file"]
        if sha256(path) != manifest[did]["sha256"]:
            fails.append(f"{did}: sha256 differs from manifest")
        if fmt == "pdf":
            text = norm(" ".join(p.extract_text() or "" for p in PdfReader(str(path)).pages))
            for q in quotes.get(did, []) + [doc["text"]]:
                if norm(q) not in text:
                    fails.append(f"{did}: pdf text layer missing {q[:50]!r}")
        elif fmt == "scan_pdf":
            text = "".join(p.extract_text() or "" for p in PdfReader(str(path)).pages).strip()
            if len(text) > 0:
                fails.append(f"{did}: scan_pdf has a text layer ({len(text)} chars)")
        elif fmt in ("photo_jpg", "handwritten_jpg", "png", "tiff_fax"):
            with Image.open(path) as im:
                im.load()
                if fmt == "tiff_fax" and (im.n_frames != 2 or im.mode != "1"):
                    fails.append(f"{did}: tiff expected 2 bit-level pages, got {im.n_frames} / {im.mode}")
        elif fmt == "docx":
            body = norm("\n".join(p.text for p in Document(str(path)).paragraphs))
            for line in doc["text"].split("\n"):
                if line.strip() and norm(line) not in body:
                    fails.append(f"{did}: docx missing line {line[:50]!r}")
        elif fmt == "txt" and path.read_bytes().decode("utf-8") != doc["text"]:
            fails.append(f"{did}: txt differs from scenario")
    fails += _verify_structured() + _verify_samples()
    return fails


def _verify_structured() -> list[str]:
    fails = []
    for sf in sc.STRUCTURED_FILES:
        path = STRUCT_DIR / sf["file"]
        if sf["format"] == "csv":
            with path.open(newline="", encoding="utf-8") as fh:
                got = list(csv.reader(fh))
            want = [sf["columns"]] + [r["values"] for r in sf["rows"]]
        else:
            ws = load_workbook(path).active
            got = [[("" if c is None else c) for c in row] for row in ws.iter_rows(values_only=True)]
            want = [sf["columns"]] + [[("" if v is None else v) for v in r["values"]] for r in sf["rows"]]
            for row in got[1:]:
                if not isinstance(row[1], (date, datetime)):
                    fails.append(f"{sf['file']}: Visit Date is not a real Excel date: {row[1]!r}")
            got = [[(c.date().isoformat() if isinstance(c, datetime) else c) for c in row] for row in got]
        if got != want:
            fails.append(f"{sf['file']}: content differs from scenario")
    return fails


def _verify_samples() -> list[str]:
    import email
    fails = []
    mixed = PdfReader(str(SAMPLES_DIR / "mixed_digital_and_scanned.pdf"))
    if len(mixed.pages) != 2 or "egfr 91" not in norm(mixed.pages[0].extract_text() or "") \
            or (mixed.pages[1].extract_text() or "").strip():
        fails.append("mixed pdf: expected text on page 1 and none on page 2")
    with zipfile.ZipFile(SAMPLES_DIR / "portal_upload_bundle.zip") as z:
        if sorted(z.namelist()) != ["D-1102.png", "D-502.jpg"]:
            fails.append("zip: wrong members")
    msg = email.message_from_bytes((SAMPLES_DIR / "email_with_attachment.eml").read_bytes())
    atts = [p.get_filename() for p in msg.walk() if p.get_filename()]
    if atts != ["D-901.pdf"]:
        fails.append(f"eml: attachments {atts}")
    if not (SAMPLES_DIR / "spoofed_extension.pdf").read_bytes().startswith(b"\xff\xd8\xff"):
        fails.append("spoofed_extension.pdf is not JPEG bytes")
    if sha256(SAMPLES_DIR / "duplicate_copy_of_D-1003.pdf") != sha256(DOCS_DIR / "D-1003.pdf"):
        fails.append("duplicate copy is not byte-identical")
    return fails


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--verify", action="store_true", help="check outputs against the scenario after generating")
    args = ap.parse_args()
    manifest = build_documents()
    build_samples()
    build_structured()
    print(f"wrote {len(manifest)} documents, samples and structured files under {DOCS_DIR.parent}")
    if args.verify:
        fails = verify()
        for f in fails:
            print("FAIL", f)
        print("verification:", "FAILED" if fails else "OK")
        return 1 if fails else 0
    return 0


if __name__ == "__main__":
    sys.exit(main())
