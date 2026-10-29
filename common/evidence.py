from __future__ import annotations

import textwrap
import uuid
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

MOCK_DISCLAIMER = (
    "MOCK EVIDENCE - FOR POC / TESTING ONLY - NOT A REAL CLINICAL RECORD"
)

SUPPORTED_FORMATS = {"pdf", "docx", "png"}
CONTENT_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "png": "image/png",
}


@dataclass(frozen=True)
class SupportingEvidence:
    source_type: str
    source_date: str
    fact_key: str
    fact_value: str
    provenance: str


@dataclass(frozen=True)
class EvidenceData:
    document_id: str
    generated_at: str
    member: dict[str, Any]
    care_gap: dict[str, Any]
    measure: dict[str, Any]
    supporting_evidence: list[SupportingEvidence]
    source_summary_points: list[str]
    generation_status: str
    review_status: str


def prepare_evidence_data(
    member: dict[str, Any],
    care_gap: dict[str, Any],
    measure: dict[str, Any],
    supporting_rows: list[dict[str, Any]],
) -> EvidenceData:
    def clean(value: Any) -> str:
        if value is None or value == "":
            return "Not available in source data"
        return str(value)

    supporting = [
        SupportingEvidence(
            source_type=clean(r.get("source_type")),
            source_date=clean(r.get("source_date")),
            fact_key=clean(r.get("fact_key")),
            fact_value=clean(r.get("fact_value")),
            provenance=clean(r.get("provenance")),
        )
        for r in supporting_rows
    ]

    if not supporting:
        supporting = [
            SupportingEvidence(
                source_type="Not available in source data",
                source_date="Not available in source data",
                fact_key="supporting_records",
                fact_value="Not available in source data",
                provenance="Not available in source data",
            )
        ]

    source_summary_points = []
    for item in supporting[:6]:
        source_summary_points.append(
            f"{item.source_type} on {item.source_date}: {item.fact_key} = {item.fact_value}"
        )

    if not source_summary_points:
        source_summary_points = ["Not available in source data"]

    return EvidenceData(
        document_id=f"EVD-{uuid.uuid4().hex[:12].upper()}",
        generated_at=datetime.now(timezone.utc).isoformat(),
        member={
            "member_id": clean(member.get("member_id")),
            "date_of_birth": clean(member.get("date_of_birth")),
            "gender": clean(member.get("gender")),
        },
        care_gap={
            "care_gap_id": clean(care_gap.get("care_gap_id")),
            "description": clean(care_gap.get("description")),
            "status": clean(care_gap.get("status")),
            "date_identified": clean(care_gap.get("date_identified")),
            "measurement_year": clean(care_gap.get("measurement_year")),
        },
        measure={
            "measure_id": clean(measure.get("measure_id")),
            "name": clean(measure.get("name")),
            "description": clean(measure.get("description")),
        },
        supporting_evidence=supporting,
        source_summary_points=source_summary_points,
        generation_status="Generated",
        review_status="Pending Review",
    )


def _header_rows(data: EvidenceData) -> list[list[str]]:
    return [
        ["Evidence Document ID", data.document_id],
        ["Generated At", data.generated_at],
    ]


def _member_rows(data: EvidenceData) -> list[list[str]]:
    return [
        ["Member ID", data.member["member_id"]],
        ["Date of Birth", data.member["date_of_birth"]],
        ["Gender", data.member["gender"]],
    ]


def _care_gap_rows(data: EvidenceData) -> list[list[str]]:
    return [
        ["Care Gap ID", data.care_gap["care_gap_id"]],
        ["Description", data.care_gap["description"]],
        ["Status", data.care_gap["status"]],
        ["Date Identified", data.care_gap["date_identified"]],
        ["Measurement Year", data.care_gap["measurement_year"]],
        ["Measure ID", data.measure["measure_id"]],
        ["Measure Name", data.measure["name"]],
    ]


def _validation_rows(data: EvidenceData) -> list[list[str]]:
    return [
        ["Evidence Generation Status", data.generation_status],
        ["Review Status", data.review_status],
        [
            "Clinical Note",
            "Generation of this document does not establish clinical validity and does not close the care gap.",
        ],
    ]


def _table_style() -> TableStyle:
    return TableStyle(
        [
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#6B7280")),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E5E7EB")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]
    )


def generate_pdf(path: Path, data: EvidenceData) -> None:
    styles = getSampleStyleSheet()
    title_style = styles["Title"]
    h2 = styles["Heading2"]
    body = ParagraphStyle(
        "Body",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
    )

    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=20 * mm,
        pageCompression=0,
    )

    elements = [
        Paragraph("HEDIS Evidence Report", title_style),
        Spacer(1, 6),
        Paragraph(f"<b>{MOCK_DISCLAIMER}</b>", body),
        Spacer(1, 10),
        Paragraph("Document Header", h2),
    ]

    for section_name, rows in [
        ("", _header_rows(data)),
        ("Member Information", _member_rows(data)),
        ("Care-Gap Information", _care_gap_rows(data)),
    ]:
        if section_name:
            elements.append(Spacer(1, 8))
            elements.append(Paragraph(section_name, h2))
        t = Table(rows, colWidths=[55 * mm, 110 * mm])
        t.setStyle(_table_style())
        elements.append(t)

    elements.extend([Spacer(1, 10), Paragraph("Supporting Evidence", h2)])
    ev_rows = [["Source Type", "Source Date", "Fact", "Value", "Provenance"]]
    for e in data.supporting_evidence:
        ev_rows.append(
            [e.source_type, e.source_date, e.fact_key, e.fact_value, e.provenance]
        )

    ev_table = Table(ev_rows, colWidths=[26 * mm, 24 * mm, 26 * mm, 58 * mm, 31 * mm])
    ev_table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#6B7280")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DBEAFE")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
            ]
        )
    )
    elements.append(ev_table)

    elements.extend([Spacer(1, 10), Paragraph("Validation and Review", h2)])
    review_table = Table(_validation_rows(data), colWidths=[55 * mm, 110 * mm])
    review_table.setStyle(_table_style())
    elements.append(review_table)

    def add_footer(canvas_obj, document):
        canvas_obj.saveState()
        footer_text = (
            f"Mock evidence for POC/testing only | {data.document_id} | "
            f"Page {document.page}"
        )
        canvas_obj.setFont("Helvetica", 8)
        canvas_obj.drawString(18 * mm, 10 * mm, footer_text)
        canvas_obj.restoreState()

    doc.build(elements, onFirstPage=add_footer, onLaterPages=add_footer)


def generate_docx(path: Path, data: EvidenceData) -> None:
    document = Document()

    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    header_title = document.add_paragraph()
    header_title.style = document.styles["Title"]
    header_title.add_run("HEDIS Evidence Report")

    subtitle = document.add_paragraph("Gap-Aligned Mock Evidence Package")
    subtitle.runs[0].italic = True

    disclaimer = document.add_paragraph()
    disclaimer.add_run(MOCK_DISCLAIMER).bold = True

    document.add_paragraph(" ")

    details = document.add_table(rows=2, cols=2)
    details.style = "Table Grid"
    details.cell(0, 0).text = "Evidence Document ID"
    details.cell(0, 1).text = data.document_id
    details.cell(1, 0).text = "Generated At"
    details.cell(1, 1).text = data.generated_at

    document.add_heading("Member Snapshot", level=2)
    _add_kv_table(document, _member_rows(data))

    document.add_heading("Care-Gap Context", level=2)
    _add_kv_table(document, _care_gap_rows(data))

    document.add_heading("Source-Derived Highlights", level=2)
    for point in data.source_summary_points:
        paragraph = document.add_paragraph(point)
        paragraph.style = "List Bullet"

    document.add_heading("Supporting Evidence", level=2)
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["Source Type", "Source Date", "Fact", "Value", "Provenance"]
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header

    for row in data.supporting_evidence:
        cells = table.add_row().cells
        cells[0].text = row.source_type
        cells[1].text = row.source_date
        cells[2].text = row.fact_key
        cells[3].text = row.fact_value
        cells[4].text = row.provenance

    document.add_heading("Validation and Review", level=2)
    _add_kv_table(document, _validation_rows(data))

    footer = document.sections[0].footer.paragraphs[0]
    footer.text = f"Mock evidence for POC/testing only | {data.document_id}"

    document.save(path)


def _add_kv_table(document: Document, rows: list[list[str]]) -> None:
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value


def _load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        "arialbd.ttf" if bold else "arial.ttf",
        "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf",
    ]
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size=size)
        except OSError:
            continue
    return ImageFont.load_default()


def _load_handwriting_font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        "segoesc.ttf",
        "comic.ttf",
        "comicbd.ttf",
        "BRADHITC.TTF",
        "Gabriola.ttf",
        "KunstlerScript.ttf",
        "DejaVuSans.ttf",
    ]
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size=size)
        except OSError:
            continue
    return ImageFont.load_default()


def _draw_handwritten_text(
    draw: ImageDraw.ImageDraw,
    text: str,
    font: ImageFont.ImageFont,
    x: int,
    y: int,
    max_width: int,
    line_spacing: int = 8,
    fill: str = "#1E3A8A",
) -> int:
    lines: list[str] = []
    words = text.split(" ")
    if not words:
        return y
    current = words[0]
    for word in words[1:]:
        trial = f"{current} {word}"
        left, _, right, _ = draw.textbbox((0, 0), trial, font=font)
        if right - left <= max_width:
            current = trial
        else:
            lines.append(current)
            current = word
    lines.append(current)

    for idx, line in enumerate(lines):
        jitter_x = 2 if idx % 2 == 0 else -1
        jitter_y = 1 if idx % 3 == 0 else 0
        draw.text((x + jitter_x, y + jitter_y), line, font=font, fill=fill)
        _, top, _, bottom = draw.textbbox((x, y), line or " ", font=font)
        y += (bottom - top) + line_spacing
    return y


def _draw_wrapped_text(
    draw: ImageDraw.ImageDraw,
    text: str,
    font: ImageFont.ImageFont,
    x: int,
    y: int,
    max_width: int,
    line_spacing: int = 4,
    fill: str = "#111827",
) -> int:
    lines: list[str] = []
    for paragraph in text.splitlines() or [""]:
        chunk = paragraph.strip() or ""
        if not chunk:
            lines.append("")
            continue
        words = chunk.split(" ")
        current = words[0]
        for word in words[1:]:
            attempt = f"{current} {word}"
            left, _, right, _ = draw.textbbox((0, 0), attempt, font=font)
            if right - left <= max_width:
                current = attempt
            else:
                lines.append(current)
                current = word
        lines.append(current)

    for line in lines:
        draw.text((x, y), line, font=font, fill=fill)
        _, top, _, bottom = draw.textbbox((x, y), line or " ", font=font)
        y += (bottom - top) + line_spacing

    return y


def generate_png(path: Path, data: EvidenceData) -> None:
    width, height = 1600, 2000
    image = Image.new("RGB", (width, height), color="#FCFCF8")
    draw = ImageDraw.Draw(image)

    # Notebook-like paper lines for handwritten medical note style.
    for y_line in range(210, height - 120, 58):
        draw.line((80, y_line, width - 80, y_line), fill="#E7E5E4", width=1)

    title_font = _load_handwriting_font(58)
    h_font = _load_handwriting_font(38)
    _load_handwriting_font(30)
    small_font = _load_handwriting_font(27)
    tiny_font = _load_handwriting_font(22)

    draw.text((90, 58), "City General Outpatient Note", font=h_font, fill="#111827")
    draw.text((92, 112), "Rx / Follow-up Observation Sheet", font=small_font, fill="#1F2937")
    draw.line((86, 168, width - 86, 168), fill="#6B7280", width=2)

    draw.text((95, 182), "HEDIS Evidence Report", font=title_font, fill="#1F2937")
    _draw_handwritten_text(
        draw,
        MOCK_DISCLAIMER,
        tiny_font,
        95,
        258,
        width - 170,
        line_spacing=5,
        fill="#991B1B",
    )

    y = 345
    draw.text((95, y), "Pt ID:", font=small_font, fill="#111827")
    _draw_handwritten_text(draw, data.member["member_id"], small_font, 250, y, 360, fill="#1E3A8A")
    draw.text((640, y), "DOB:", font=small_font, fill="#111827")
    _draw_handwritten_text(draw, data.member["date_of_birth"], small_font, 760, y, 300, fill="#1E3A8A")
    draw.text((1110, y), "Sex:", font=small_font, fill="#111827")
    _draw_handwritten_text(draw, data.member["gender"], small_font, 1220, y, 250, fill="#1E3A8A")

    y += 72
    draw.text((95, y), "Care Gap:", font=small_font, fill="#111827")
    _draw_handwritten_text(draw, data.care_gap["care_gap_id"], small_font, 285, y, 280, fill="#1E3A8A")
    draw.text((610, y), "Status:", font=small_font, fill="#111827")
    _draw_handwritten_text(draw, data.care_gap["status"], small_font, 770, y, 240, fill="#1E3A8A")
    draw.text((1040, y), "Year:", font=small_font, fill="#111827")
    _draw_handwritten_text(draw, data.care_gap["measurement_year"], small_font, 1185, y, 220, fill="#1E3A8A")

    y += 70
    draw.text((95, y), "Measure:", font=small_font, fill="#111827")
    _draw_handwritten_text(
        draw,
        f"{data.measure['measure_id']} - {data.measure['name']}",
        small_font,
        290,
        y,
        width - 380,
        fill="#1E3A8A",
    )

    y += 72
    draw.text((95, y), "Assessment:", font=small_font, fill="#111827")
    y = _draw_handwritten_text(
        draw,
        data.care_gap["description"],
        small_font,
        300,
        y,
        width - 390,
        line_spacing=8,
        fill="#1E3A8A",
    )

    y += 34
    draw.text((95, y), "Source Notes:", font=small_font, fill="#111827")
    y += 48
    for row in data.supporting_evidence[:8]:
        evidence_line = (
            f"- {row.source_date} {row.source_type}: {row.fact_key} = {row.fact_value}"
        )
        y_text = _draw_handwritten_text(
            draw,
            evidence_line,
            small_font,
            130,
            y,
            width - 180,
            line_spacing=5,
            fill="#1E3A8A",
        )
        y = y_text + 10

    if y < height - 260:
        y += 14
    draw.text((95, y), "Validation:", font=small_font, fill="#111827")
    _draw_handwritten_text(
        draw,
        f"{data.generation_status} / {data.review_status}",
        small_font,
        290,
        y,
        520,
        fill="#1E3A8A",
    )

    stamp_x1, stamp_y1, stamp_x2, stamp_y2 = width - 470, height - 260, width - 140, height - 120
    draw.ellipse((stamp_x1, stamp_y1, stamp_x2, stamp_y2), outline="#B91C1C", width=4)
    draw.text((stamp_x1 + 48, stamp_y1 + 38), "POC MOCK", font=small_font, fill="#B91C1C")
    draw.text((stamp_x1 + 84, stamp_y1 + 82), "EVIDENCE", font=tiny_font, fill="#B91C1C")

    draw.line((86, height - 88, width - 86, height - 88), fill="#6B7280", width=1)
    draw.text(
        (95, height - 74),
        f"Doc: {data.document_id}   Generated: {data.generated_at}",
        font=tiny_font,
        fill="#374151",
    )

    image.save(path, format="PNG")


def extension_for_format(evidence_format: str) -> str:
    return {
        "pdf": ".pdf",
        "docx": ".docx",
        "png": ".png",
    }[evidence_format]


def generate_document(evidence_format: str, path: Path, data: EvidenceData) -> None:
    if evidence_format == "pdf":
        generate_pdf(path, data)
    elif evidence_format == "docx":
        generate_docx(path, data)
    elif evidence_format == "png":
        generate_png(path, data)
    else:
        raise ValueError(f"Unsupported format: {evidence_format}")


def ensure_supported_formats(formats: list[str]) -> list[str]:
    normalized = [f.lower() for f in formats]
    if not normalized:
        raise ValueError("At least one output format must be requested")
    invalid = [f for f in normalized if f not in SUPPORTED_FORMATS]
    if invalid:
        raise ValueError(f"Unsupported format(s): {', '.join(invalid)}")
    return normalized


def safe_text(value: str, max_len: int = 120) -> str:
    return textwrap.shorten(value, width=max_len, placeholder="...")
