"""CSV and Excel import with column mapping (US-02 T4).

A structured file (lab export, vitals workbook) is read into rows, matched to a saved mapping by its
exact header list, and every row is validated. Each data row becomes one structured_records row:

  loaded           a valid, coded record
  rejected         a row that failed a check, with the reason in plain words
  free_text_to_ai  a row whose only content is a free-text column; it also gets a small document so the
                   AI extraction step can read it like any other note

A file whose layout has no mapping is not guessed at: the headers and a keyword-based suggestion are
returned so a person can confirm a mapping (save_mapping), then the file is loaded again.
"""

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import psycopg
import yaml
from openpyxl import load_workbook
from psycopg.types.json import Jsonb

from .db import audit
from .document_service import save_conversion
from .models import PLAIN_TEXT, ConversionResult, Page

MAPPINGS_DIR = Path(__file__).resolve().parent / "mappings"

RECORD_TYPES = ("lab_result", "vital_sign")
FIELDS = ("member_id", "member_name", "birth_date", "test_name", "value", "unit", "observed_on",
          "systolic", "diastolic", "pulse", "notes", "provider")
REQUIRED_FIELDS = {
    "lab_result": ("member_id", "test_name", "value", "unit", "observed_on"),
    "vital_sign": ("member_id", "systolic", "diastolic", "observed_on"),
}
# Fields that carry clinical content. A row with none of these but some free text is a free-text row.
CLINICAL_FIELDS = {
    "lab_result": ("test_name", "value", "unit"),
    "vital_sign": ("systolic", "diastolic", "pulse"),
}
DATE_LABEL = {"lab_result": "collection date", "vital_sign": "visit date"}

BP_CODE, BP_SYSTEM, BP_DESCRIPTION, BP_UNIT = "85354-9", "LOINC", "Blood pressure panel", "mm[Hg]"
FREE_TEXT_DOC_TYPE = "structured free text"
UNKNOWN_LAYOUT_NOTE = "unknown layout: confirm a mapping"

COMMON_DATE_FORMATS = ["%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%m-%d-%Y", "%Y/%m/%d", "%m/%d/%y"]
DEFAULT_DATE_FORMATS = ["%Y-%m-%d", "%m/%d/%Y"]
ADVISORY_KEYS = ("unmapped_headers", "notes")   # returned in a suggestion, never saved

Row = dict[str, Any]


class MappingError(ValueError):
    """A mapping is not valid. `problems` lists every reason in plain words."""

    def __init__(self, problems: list[str]) -> None:
        super().__init__("; ".join(problems))
        self.problems = problems


# ---------------------------------------------------------------------------------------------
# Reading files
# ---------------------------------------------------------------------------------------------

def _clean(value: Any) -> Any:
    """Empty cells become None; text is trimmed; Excel dates become ISO date text."""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _headers(cells: list[Any]) -> list[str]:
    """Header row text; trailing empty header cells are dropped, other blanks are an error."""
    names = [("" if c is None else str(c).strip()) for c in cells]
    while names and not names[-1]:
        names.pop()
    if not names:
        raise ValueError("the file has no header row")
    if not all(names):
        raise ValueError("the header row has an empty column name")
    if len(set(names)) != len(names):
        raise ValueError("the header row has a repeated column name")
    return names


def _read_csv(data: bytes) -> tuple[list[str], list[Row]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("the file is not UTF-8 text") from exc
    reader = csv.reader(io.StringIO(text, newline=""))
    first = next(reader, None)
    if first is None:
        raise ValueError("the file is empty")
    headers = _headers(first)
    rows: list[Row] = []
    for cells in reader:
        values = [_clean(c) for c in cells[:len(headers)]]
        values += [None] * (len(headers) - len(values))
        if any(v is not None for v in values):
            rows.append(dict(zip(headers, values)))
    return headers, rows


def _read_xlsx(data: bytes) -> tuple[list[str], list[Row]]:
    try:
        book = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError("the file is not a readable Excel workbook") from exc
    try:
        lines = book.worksheets[0].iter_rows(values_only=True)
        first = next(lines, None)
        if first is None:
            raise ValueError("the first sheet is empty")
        headers = _headers(list(first))
        rows: list[Row] = []
        for cells in lines:
            values = [_clean(c) for c in list(cells)[:len(headers)]]
            values += [None] * (len(headers) - len(values))
            if any(v is not None for v in values):
                rows.append(dict(zip(headers, values)))
    finally:
        book.close()
    return headers, rows


def read_rows(data: bytes, filename: str) -> tuple[list[str], list[Row]]:
    """Read a CSV or XLSX file into (headers, rows).

    CSV is read as UTF-8 (a leading byte-order mark is fine). XLSX uses the first sheet with headers in
    row 1; real Excel dates become ISO date text. Empty cells are None and completely blank rows are
    skipped, so the Nth returned row is the Nth data row of the file.
    """
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return _read_csv(data)
    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx(data)
    raise ValueError(f"unsupported structured file type '{suffix or filename}'; use .csv or .xlsx")


# ---------------------------------------------------------------------------------------------
# Mappings
# ---------------------------------------------------------------------------------------------

def _norm(text: str) -> str:
    """Lower-case with single spaces; used to compare test names."""
    return re.sub(r"\s+", " ", text).strip().lower()


def _norm_unit(text: str) -> str:
    return re.sub(r"\s+", "", text).lower()


@dataclass(frozen=True)
class Mapping:
    """One source layout: how its columns become structured record fields."""
    name: str
    version: int
    match: tuple[str, ...]
    record_type: str
    columns: dict[str, str]                 # header -> field
    date_formats: tuple[str, ...]
    value_maps: dict[str, dict[str, dict]]  # field -> source value -> coding (only test_name is used)
    free_text_columns: tuple[str, ...] = ()
    _tests: dict[str, dict] = field(default_factory=dict, repr=False, compare=False)

    def header_for(self, field_name: str) -> str | None:
        """The file header mapped to a field, if any."""
        return next((h for h, f in self.columns.items() if f == field_name), None)

    def lookup_test(self, test_name: str) -> dict | None:
        """The coding entry for a test name as written in the file, ignoring case and spacing."""
        return self._tests.get(_norm(test_name))

    def to_dict(self) -> dict:
        """The plain dict form that is written to YAML."""
        return {
            "name": self.name, "version": self.version, "record_type": self.record_type,
            "match": list(self.match), "columns": dict(self.columns),
            "date_formats": list(self.date_formats), "value_maps": self.value_maps,
            "free_text_columns": list(self.free_text_columns),
        }


def validate_mapping(raw: Any) -> list[str]:
    """Every problem found in a mapping dict, in plain words. An empty list means it is valid."""
    if not isinstance(raw, dict):
        return ["a mapping must be an object with name, match, record_type and columns"]
    problems: list[str] = []
    name = raw.get("name")
    if not isinstance(name, str) or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]*", name):
        problems.append("name is required and may only use letters, digits, underscore and hyphen")
    if not isinstance(raw.get("version"), int) or isinstance(raw.get("version"), bool) or raw["version"] < 1:
        problems.append("version must be a whole number of 1 or more")
    record_type = raw.get("record_type")
    if record_type not in RECORD_TYPES:
        problems.append(f"record_type must be one of {', '.join(RECORD_TYPES)}")

    match = raw.get("match")
    match_ok = (isinstance(match, list) and bool(match) and all(isinstance(h, str) and h.strip() for h in match)
                and len(set(match)) == len(match))
    if not match_ok:
        problems.append("match must be a list of the file's header names, with no blanks or repeats")
        match = []

    columns = raw.get("columns")
    columns_ok = isinstance(columns, dict) and all(isinstance(k, str) and isinstance(v, str)
                                                   for k, v in columns.items())
    if not columns_ok:
        problems.append("columns must map each header to a field")
        columns = {}
    problems += _column_problems(columns, match, record_type)
    problems += _format_problems(raw.get("date_formats"))
    problems += _value_map_problems(raw.get("value_maps"), record_type)
    problems += _free_text_problems(raw.get("free_text_columns"), match)
    return problems


def _column_problems(columns: dict, match: list, record_type: Any) -> list[str]:
    problems = []
    for header, field_name in columns.items():
        if header not in match:
            problems.append(f"column '{header}' is not in the match header list")
        if field_name not in FIELDS:
            problems.append(f"column '{header}' maps to unknown field '{field_name}'")
    mapped = list(columns.values())
    for field_name in sorted({f for f in mapped if mapped.count(f) > 1}):
        problems.append(f"field {field_name} is mapped from more than one column")
    for field_name in REQUIRED_FIELDS.get(record_type, ()):
        if field_name not in mapped:
            problems.append(f"{field_name} is not mapped; a {record_type} needs it")
    return problems


def _format_problems(formats: Any) -> list[str]:
    if not isinstance(formats, list) or not formats or not all(isinstance(f, str) for f in formats):
        return ["date_formats must list at least one date format such as %m/%d/%Y"]
    problems = []
    probe = datetime(2001, 2, 3)
    for fmt in formats:
        try:
            datetime.strptime(probe.strftime(fmt), fmt)
        except ValueError:
            problems.append(f"date format '{fmt}' is not valid")
    return problems


def _value_map_problems(value_maps: Any, record_type: Any) -> list[str]:
    if not isinstance(value_maps, dict):
        return ["value_maps must be an object (use {} when there is nothing to map)"]
    problems = [f"value_maps.{key} is not supported; only test_name is" for key in value_maps if key != "test_name"]
    tests = value_maps.get("test_name")
    if record_type == "lab_result" and not tests:
        problems.append("value_maps.test_name must list each test name with its code, so results can be coded")
    for alias, entry in (tests or {}).items():
        ok = (isinstance(entry, dict)
              and all(isinstance(entry.get(k), str) and entry[k].strip() for k in ("code", "code_system", "description"))
              and isinstance(entry.get("allowed_units"), list) and bool(entry["allowed_units"])
              and all(isinstance(u, str) and u.strip() for u in entry["allowed_units"]))
        if not ok:
            problems.append(f"test '{alias}' needs code, code_system, description and a list of allowed_units")
    return problems


def _free_text_problems(free_text: Any, match: list) -> list[str]:
    if not isinstance(free_text, list) or not all(isinstance(h, str) for h in free_text):
        return ["free_text_columns must be a list of header names (it may be empty)"]
    return [f"free text column '{h}' is not in the match header list" for h in free_text if h not in match]


def parse_mapping(raw: Any) -> Mapping:
    """Build a Mapping from a dict, raising MappingError listing every problem."""
    problems = validate_mapping({k: v for k, v in raw.items() if k not in ADVISORY_KEYS}
                                if isinstance(raw, dict) else raw)
    if problems:
        raise MappingError(problems)
    tests = {_norm(alias): entry for alias, entry in (raw["value_maps"].get("test_name") or {}).items()}
    return Mapping(
        name=raw["name"], version=raw["version"], match=tuple(raw["match"]), record_type=raw["record_type"],
        columns=dict(raw["columns"]), date_formats=tuple(raw["date_formats"]),
        value_maps=raw["value_maps"], free_text_columns=tuple(raw["free_text_columns"]), _tests=tests)


def load_mappings() -> list[Mapping]:
    """Every saved mapping, in name order."""
    found = []
    for path in sorted(MAPPINGS_DIR.glob("*.yaml")):
        try:
            found.append(parse_mapping(yaml.safe_load(path.read_text(encoding="utf-8"))))
        except (MappingError, yaml.YAMLError) as exc:
            raise MappingError([f"mapping file {path.name} is not valid: {exc}"]) from exc
    return found


def find_mapping(headers: list[str]) -> Mapping | None:
    """The saved mapping whose header list equals these headers exactly (order does not matter)."""
    wanted = sorted(h.strip() for h in headers)
    for mapping in load_mappings():
        if sorted(mapping.match) == wanted:
            return mapping
    return None


def save_mapping(mapping: dict) -> Path:
    """Validate a confirmed mapping and write it as <name>.yaml. Returns the file path.

    Refuses a mapping that is invalid, that reuses an existing name, or whose headers another mapping
    already recognises (two mappings for one layout would be ambiguous).
    """
    parsed = parse_mapping(mapping)
    path = MAPPINGS_DIR / f"{parsed.name}.yaml"
    if path.exists():
        raise MappingError([f"a mapping named '{parsed.name}' already exists"])
    clash = find_mapping(list(parsed.match))
    if clash is not None:
        raise MappingError([f"mapping '{clash.name}' already recognises these headers"])
    MAPPINGS_DIR.mkdir(parents=True, exist_ok=True)
    path.write_text(yaml.dump(parsed.to_dict(), Dumper=_NoAliasDumper, sort_keys=False, allow_unicode=True),
                    encoding="utf-8")
    return path


class _NoAliasDumper(yaml.SafeDumper):
    """Writes shared entries in full instead of using YAML anchors."""

    def ignore_aliases(self, data: Any) -> bool:
        return True


# ---------------------------------------------------------------------------------------------
# Suggesting a mapping for an unknown layout (keyword rules, no AI)
# ---------------------------------------------------------------------------------------------

# First rule that matches a header wins; a field is taken by the first header that matches it.
_SUGGEST_RULES: list[tuple[str, str]] = [
    ("birth_date", r"\b(dob|birth|born)\b"),
    ("member_id", r"\b(member|patient|subscriber|mrn)\b.*\b(id|number|no)\b|^(mrn|member|patient)$|^(member|patient) ?id$"),
    ("member_name", r"\b(member|patient)\b.*\bname\b|^name$"),
    ("systolic", r"\b(systolic|sbp)\b"),
    ("diastolic", r"\b(diastolic|dbp)\b"),
    ("pulse", r"\b(pulse|heart rate|hr)\b"),
    ("test_name", r"\b(test|analyte|component|lab name|procedure)\b"),
    ("unit", r"\bunits?\b|\buom\b"),
    ("value", r"\b(result|value|reading)\b"),
    ("observed_on", r"\b(collected|collection|visit|service|observation|observed|performed|draw|dos)\b|\bdate\b"),
    ("notes", r"\b(note|notes|comment|comments|remark|remarks|narrative)\b"),
    ("provider", r"\b(provider|physician|doctor|ordering|clinician)\b"),
]


def _slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")


def _infer_date_formats(samples: list[str]) -> list[str]:
    """Date formats (from common ones) that read every sample value; a default when there are no samples."""
    if not samples:
        return list(DEFAULT_DATE_FORMATS)
    fits = []
    for fmt in COMMON_DATE_FORMATS:
        try:
            for s in samples:
                datetime.strptime(s, fmt)
        except ValueError:
            continue
        fits.append(fmt)
    return fits or list(DEFAULT_DATE_FORMATS)


def suggest_mapping(headers: list[str], sample_rows: list[Row] | None = None, name_hint: str = "") -> dict:
    """A best-guess mapping for a layout nobody has mapped yet, for a person to review and confirm.

    Uses keyword rules on the header names only (and sample rows, if given, to guess date formats).
    It never invents test codes: value_maps stays empty for a person to fill in. The result also carries
    `unmapped_headers` and `notes`, which are advice for the reviewer and are not saved.
    """
    columns: dict[str, str] = {}
    taken: set[str] = set()
    for header in headers:
        text = re.sub(r"[^a-z0-9]+", " ", header.lower()).strip()
        for field_name, pattern in _SUGGEST_RULES:
            if field_name not in taken and re.search(pattern, text):
                columns[header] = field_name
                taken.add(field_name)
                break

    if "systolic" in taken or "diastolic" in taken:
        record_type = "vital_sign"
    elif "test_name" in taken or "value" in taken:
        record_type = "lab_result"
    else:
        record_type = None

    unmapped = [h for h in headers if h not in columns]
    notes = ["Confirm every column, add the test codes under value_maps.test_name, then save the mapping."]
    if record_type is None:
        notes.append("Could not tell if this is lab results or vital signs: set record_type.")
    missing = [f for f in REQUIRED_FIELDS.get(record_type or "", ()) if f not in taken]
    if missing:
        notes.append("No column found for: " + ", ".join(missing) + ".")

    date_headers = [h for h, f in columns.items() if f in ("observed_on", "birth_date")]
    samples = [str(r[h]) for r in (sample_rows or []) for h in date_headers if r.get(h) is not None]
    return {
        "name": _slug(name_hint),
        "version": 1,
        "record_type": record_type,
        "match": list(headers),
        "columns": columns,
        "date_formats": _infer_date_formats(samples),
        "value_maps": {},
        "free_text_columns": [h for h, f in columns.items() if f == "notes"],
        "unmapped_headers": unmapped,
        "notes": notes,
    }


# ---------------------------------------------------------------------------------------------
# Validating rows
# ---------------------------------------------------------------------------------------------

@dataclass
class _Outcome:
    """The result of checking one row."""
    status: str                                   # loaded | rejected | free_text_to_ai
    fields: dict[str, Any] = field(default_factory=dict)   # structured_records columns to store
    reasons: list[str] = field(default_factory=list)
    free_text: str | None = None


def _to_decimal(value: Any) -> Decimal:
    try:
        number = Decimal(str(value).strip())
    except InvalidOperation as exc:
        raise ValueError(f"'{value}' is not a number") from exc
    if not number.is_finite():
        raise ValueError(f"'{value}' is not a number")
    return number


def _to_date(value: Any, formats: tuple[str, ...]) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"'{text}' is not a valid date")


def _check_member_and_date(values: Row, mapping: Mapping, members: dict[str, str], out: _Outcome) -> None:
    member = values.get("member_id")
    if member is None:
        out.reasons.append("missing member id")
    else:
        member = str(member).strip()
        if member in members:
            out.fields["patient_id"] = members[member]
        else:
            out.reasons.append(f"member {member} not found")
    label = DATE_LABEL[mapping.record_type]
    when = values.get("observed_on")
    if when is None:
        out.reasons.append(f"missing {label}")
        return
    try:
        out.fields["observed_on"] = _to_date(when, mapping.date_formats)
    except ValueError:
        out.reasons.append(f"{label} '{when}' is not a valid date")


def _check_lab(values: Row, mapping: Mapping, out: _Outcome) -> None:
    test, entry = values.get("test_name"), None
    if test is None:
        out.reasons.append("missing test name")
    else:
        entry = mapping.lookup_test(str(test))
        if entry is None:
            out.reasons.append(f"test {test} is not in the mapping")
    if entry is not None:
        out.fields.update(code=entry["code"], code_system=entry["code_system"], description=entry["description"])

    result = values.get("value")
    if result is None:
        out.reasons.append("missing result")
    else:
        try:
            out.fields["value_num"] = _to_decimal(result)
            out.fields["value_text"] = str(result).strip()
        except ValueError:
            out.reasons.append(f"result '{result}' is not a number")

    unit = values.get("unit")
    if unit is None:
        out.reasons.append("missing unit")
    elif entry is not None:
        canonical = {_norm_unit(u): u for u in entry["allowed_units"]}.get(_norm_unit(str(unit)))
        if canonical is None:
            out.reasons.append(f"unit {unit} is not valid for {entry['description']}")
        else:
            out.fields["unit"] = canonical


def _check_vital(values: Row, out: _Outcome) -> None:
    readings: dict[str, int] = {}
    for key in ("systolic", "diastolic"):
        raw = values.get(key)
        if raw is None:
            out.reasons.append(f"{key} missing")
            continue
        try:
            number = _to_decimal(raw)
        except ValueError:
            out.reasons.append(f"{key} '{raw}' is not a number")
            continue
        if number != number.to_integral_value() or number <= 0:
            out.reasons.append(f"{key} '{raw}' must be a positive whole number")
        else:
            readings[key] = int(number)
    if len(readings) == 2:
        out.fields.update(value_text=f"{readings['systolic']}/{readings['diastolic']}", code=BP_CODE,
                          code_system=BP_SYSTEM, description=BP_DESCRIPTION, unit=BP_UNIT)


def _check_row(row: Row, mapping: Mapping, members: dict[str, str]) -> _Outcome:
    """Validate one row against its mapping and decide loaded, rejected or free_text_to_ai."""
    values = {field_name: row.get(header) for header, field_name in mapping.columns.items()}
    out = _Outcome("loaded", fields={"record_type": mapping.record_type})
    free_text = "\n".join(str(row[h]) for h in mapping.free_text_columns if row.get(h) is not None)
    no_clinical = all(values.get(f) is None for f in CLINICAL_FIELDS[mapping.record_type])

    _check_member_and_date(values, mapping, members, out)
    if no_clinical and free_text:
        # A free-text row needs a known member, but its date is optional context for the document.
        out.reasons = [r for r in out.reasons if r != f"missing {DATE_LABEL[mapping.record_type]}"]
        out.status, out.free_text = "free_text_to_ai", free_text
        out.fields["value_text"] = free_text
    elif mapping.record_type == "lab_result":
        _check_lab(values, mapping, out)
    else:
        _check_vital(values, out)

    if out.reasons:
        out.status = "rejected"
    return out


# ---------------------------------------------------------------------------------------------
# Loading a file
# ---------------------------------------------------------------------------------------------

def _clear_previous_load(conn: psycopg.Connection, source_file_id: int) -> dict:
    """Remove records and free-text documents from an earlier load of this file.

    Nothing may reference them yet: once AI evidence or a supplemental record points at one of the
    small documents, reloading would break that trail, so it is refused.
    """
    doc_ids = [r[0] for r in conn.execute(
        "SELECT document_reference_id FROM document_conversion "
        "WHERE source_file_id = %s AND document_reference_id LIKE %s",
        (source_file_id, f"S-{source_file_id}-%")).fetchall()]
    if doc_ids:
        for table, column in (("evidence_proposal_detail", "document_reference_id"),
                              ("supplemental_records", "document_reference_id")):
            if conn.execute(f"SELECT 1 FROM {table} WHERE {column} = ANY(%s) LIMIT 1", (doc_ids,)).fetchone():
                raise RuntimeError(f"file {source_file_id} was already used for evidence; it cannot be reloaded")
        conn.execute("UPDATE document_conversion SET near_duplicate_of = NULL WHERE near_duplicate_of = ANY(%s)",
                     (doc_ids,))
        conn.execute("DELETE FROM page_words WHERE page_id IN "
                     "(SELECT id FROM document_pages WHERE document_reference_id = ANY(%s))", (doc_ids,))
        conn.execute("DELETE FROM document_pages WHERE document_reference_id = ANY(%s)", (doc_ids,))
        conn.execute("DELETE FROM document_conversion WHERE document_reference_id = ANY(%s)", (doc_ids,))
        conn.execute("DELETE FROM document_reference WHERE document_reference_id = ANY(%s)", (doc_ids,))
    removed = conn.execute("DELETE FROM structured_records WHERE source_file_id = %s", (source_file_id,)).rowcount
    return {"records": removed, "documents": len(doc_ids)}


def _insert_record(conn: psycopg.Connection, source_file_id: int, row_number: int, row: Row, out: _Outcome) -> None:
    f = out.fields
    conn.execute(
        """INSERT INTO structured_records (source_file_id, row_number, patient_id, record_type, code_system, code,
               description, value_text, value_num, unit, observed_on, status, reject_reason, raw)
           VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)""",
        (source_file_id, row_number, f.get("patient_id"), f.get("record_type"), f.get("code_system"), f.get("code"),
         f.get("description"), f.get("value_text"), f.get("value_num"), f.get("unit"), f.get("observed_on"),
         out.status, "; ".join(out.reasons) if out.reasons else None, Jsonb(row)))


def _create_free_text_document(conn: psycopg.Connection, source_file_id: int, row_number: int, out: _Outcome) -> None:
    """A one-page text document for the AI extraction step, from a row's free-text content."""
    doc_id = f"S-{source_file_id}-{row_number}"
    observed = out.fields.get("observed_on")
    conn.execute(
        """INSERT INTO document_reference (document_reference_id, patient_id, status, type_code, doc_date)
           VALUES (%s, %s, 'current', %s, %s)""",
        (doc_id, out.fields.get("patient_id"), FREE_TEXT_DOC_TYPE,
         datetime.combine(observed, time(0), tzinfo=timezone.utc) if observed else None))
    conn.execute(
        """INSERT INTO document_conversion (document_reference_id, source_file_id, source_org_id)
           SELECT %s, id, source_org_id FROM source_files WHERE id = %s""", (doc_id, source_file_id))
    save_conversion(conn, doc_id, ConversionResult(pages=[Page(page_no=1, text=out.free_text or "",
                                                               method=PLAIN_TEXT)]))


def _fail_file(conn: psycopg.Connection, source_file_id: int, note: str) -> None:
    conn.execute("UPDATE source_files SET processing_status = 'failed', processing_note = %s WHERE id = %s",
                 (note, source_file_id))


def load_structured_file(conn: psycopg.Connection, source_file_id: int, data: bytes, filename: str) -> dict:
    """Load a CSV or XLSX file into structured_records, one row per data row.

    Returns {status: 'loaded', loaded, rejected, free_text_to_ai, errors: [{row, reason}]}.
    When no saved mapping fits the headers it changes nothing but the file's status and returns
    {status: 'needs_mapping', headers, suggestion}. When the file cannot be read at all it returns
    {status: 'failed', reason}. Loading the same file again replaces the earlier result.
    The caller owns the transaction and the file's final processing_status.
    """
    try:
        headers, rows = read_rows(data, filename)
    except ValueError as exc:
        _fail_file(conn, source_file_id, f"could not read file: {exc}")
        audit(conn, "source_file", str(source_file_id), "structured.failed", {"reason": str(exc)})
        return {"status": "failed", "reason": str(exc)}

    mapping = find_mapping(headers)
    if mapping is None:
        _fail_file(conn, source_file_id, UNKNOWN_LAYOUT_NOTE)
        audit(conn, "source_file", str(source_file_id), "structured.needs_mapping", {"headers": headers})
        return {"status": "needs_mapping", "headers": headers,
                "suggestion": suggest_mapping(headers, rows[:20], Path(filename).stem)}

    replaced = _clear_previous_load(conn, source_file_id)
    id_header = mapping.header_for("member_id")
    ids = {str(r[id_header]).strip() for r in rows if r.get(id_header) is not None}
    members = dict(conn.execute("SELECT mrn, min(patient_id) FROM patient WHERE mrn = ANY(%s) GROUP BY mrn",
                                (sorted(ids),)).fetchall())

    counts = {"loaded": 0, "rejected": 0, "free_text_to_ai": 0}
    errors: list[dict] = []
    for row_number, row in enumerate(rows, start=1):
        out = _check_row(row, mapping, members)
        _insert_record(conn, source_file_id, row_number, row, out)
        counts[out.status] += 1
        if out.status == "rejected":
            errors.append({"row": row_number, "reason": "; ".join(out.reasons)})
        elif out.status == "free_text_to_ai":
            _create_free_text_document(conn, source_file_id, row_number, out)

    note = f"structured: {counts['loaded']} loaded, {counts['rejected']} rejected, {counts['free_text_to_ai']} to AI"
    conn.execute("UPDATE source_files SET processing_note = %s WHERE id = %s", (note, source_file_id))
    audit(conn, "source_file", str(source_file_id), "structured.loaded",
          {"mapping": mapping.name, "mapping_version": mapping.version, "rows": len(rows), **counts,
           "replaced_records": replaced["records"], "replaced_documents": replaced["documents"]})
    return {"status": "loaded", **counts, "errors": errors}
